"""PDF、DOCX、TXT 简历的统一解析工具。"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

import fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


MAX_TEXT_CHARACTERS = 50_000
MAX_TXT_BYTES = 2 * 1024 * 1024
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 100
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class ResumeParseError(ValueError):
    """简历无法安全解析时抛出的用户可读异常。"""


# 保留阶段 1 的类名，避免既有调用方失效。
TextParseError = ResumeParseError


@dataclass(frozen=True, slots=True)
class ParsedResume:
    """统一解析结果，供页面展示文件格式和文本统计信息。"""

    text: str
    file_type: str
    character_count: int
    page_count: int | None = None


def _clean_extracted_text(text: str, source_label: str) -> str:
    """清理提取文本，同时保留段落结构并统一执行长度检查。"""
    normalized = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    cleaned_lines: list[str] = []
    previous_blank = False

    for raw_line in normalized.split("\n"):
        line = re.sub(r"[ \t]+", " ", raw_line).strip()
        if line:
            cleaned_lines.append(line)
            previous_blank = False
        elif not previous_blank and cleaned_lines:
            cleaned_lines.append("")
            previous_blank = True

    cleaned = "\n".join(cleaned_lines).strip()
    if not cleaned:
        raise ResumeParseError(
            f"{source_label}中未提取到可读文本；如果是扫描件，请先进行 OCR 后再上传。"
        )
    if len(cleaned) > MAX_TEXT_CHARACTERS:
        raise ResumeParseError(
            f"解析后的简历超过 {MAX_TEXT_CHARACTERS:,} 个字符，请精简内容后重新上传。"
        )
    return cleaned


def parse_txt_bytes(content: bytes, max_characters: int = MAX_TEXT_CHARACTERS) -> str:
    """将 TXT 字节解析为文本，兼容 UTF-8 和常见中文 Windows 编码。"""
    if not content:
        raise ResumeParseError("上传的 TXT 文件为空，请选择包含简历内容的文件。")
    if len(content) > MAX_TXT_BYTES:
        raise ResumeParseError("TXT 文件超过 2 MB，请精简内容后重新上传。")

    decoded_text: str | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            decoded_text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if decoded_text is None:
        raise ResumeParseError("TXT 文件编码无法识别，请将文件另存为 UTF-8 后重试。")

    cleaned_text = decoded_text.strip()
    if not cleaned_text:
        raise ResumeParseError("TXT 文件只包含空白字符，请检查文件内容。")
    if len(cleaned_text) > max_characters:
        raise ResumeParseError(
            f"简历文本超过 {max_characters:,} 个字符，请精简内容后重新上传。"
        )
    return cleaned_text


def parse_pdf_bytes(content: bytes) -> tuple[str, int]:
    """解析 PDF 文本并返回正文与页数，不对扫描图片进行隐式 OCR。"""
    if b"%PDF-" not in content[:1024]:
        raise ResumeParseError("文件扩展名是 PDF，但内容不是有效的 PDF 文件。")

    try:
        with fitz.open(stream=content, filetype="pdf") as document:
            if document.needs_pass:
                raise ResumeParseError("PDF 已加密，请先移除打开密码后再上传。")
            if document.page_count == 0:
                raise ResumeParseError("PDF 没有可解析的页面。")
            if document.page_count > MAX_PDF_PAGES:
                raise ResumeParseError(
                    f"PDF 超过 {MAX_PDF_PAGES} 页，请仅保留与简历相关的页面。"
                )

            page_texts = [page.get_text("text", sort=True) for page in document]
            page_count = document.page_count
    except ResumeParseError:
        raise
    except (fitz.FileDataError, fitz.EmptyFileError, RuntimeError, ValueError) as exc:
        raise ResumeParseError("PDF 文件损坏或格式异常，无法读取。") from exc

    return _clean_extracted_text("\n\n".join(page_texts), "PDF"), page_count


def parse_docx_bytes(content: bytes) -> str:
    """解析 DOCX 的正文段落和表格文本。"""
    if not content.startswith(b"PK"):
        raise ResumeParseError("文件扩展名是 DOCX，但内容不是有效的 DOCX 文件。")

    try:
        document = Document(io.BytesIO(content))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError) as exc:
        raise ResumeParseError("DOCX 文件损坏或格式异常，无法读取。") from exc

    text_blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                text_blocks.append("\t".join(cells))

    return _clean_extracted_text("\n".join(text_blocks), "DOCX")


def parse_resume_bytes(content: bytes, filename: str) -> ParsedResume:
    """根据文件扩展名调用对应解析器，并返回统一结果。"""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError("仅支持 PDF、DOCX 和 TXT 格式的简历文件。")
    if not content:
        raise ResumeParseError("上传的文件为空，请重新选择简历文件。")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ResumeParseError("简历文件超过 10 MB，请压缩或精简后重新上传。")

    if extension == ".txt":
        text = parse_txt_bytes(content)
        return ParsedResume(text=text, file_type="TXT", character_count=len(text))
    if extension == ".pdf":
        text, page_count = parse_pdf_bytes(content)
        return ParsedResume(
            text=text,
            file_type="PDF",
            character_count=len(text),
            page_count=page_count,
        )

    text = parse_docx_bytes(content)
    return ParsedResume(text=text, file_type="DOCX", character_count=len(text))
