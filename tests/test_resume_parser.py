"""PDF、DOCX、TXT 简历统一解析测试。"""

import io

import fitz
import pytest
from docx import Document

from utils.resume_parser import (
    MAX_DOCUMENT_BYTES,
    MAX_TXT_BYTES,
    ResumeParseError,
    parse_resume_bytes,
    parse_txt_bytes,
)


def _make_docx_bytes() -> bytes:
    """在内存中创建带正文和表格的 DOCX 测试文件。"""
    document = Document()
    document.add_paragraph("技能：Python、SQL")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "Flask 系统"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _make_pdf_bytes(text: str | None = "Python Flask SQL resume") -> bytes:
    """在内存中创建 PDF；文本为空时模拟图片型扫描简历。"""
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def test_parse_utf8_text() -> None:
    assert parse_txt_bytes("技能：Python、SQL".encode("utf-8")) == "技能：Python、SQL"


def test_parse_gb18030_text() -> None:
    assert parse_txt_bytes("中文简历".encode("gb18030")) == "中文简历"


def test_unified_txt_result() -> None:
    result = parse_resume_bytes("Python 简历".encode("utf-8"), "resume.TXT")
    assert result.file_type == "TXT"
    assert result.text == "Python 简历"
    assert result.character_count == len(result.text)
    assert result.page_count is None


def test_parse_docx_paragraph_and_table() -> None:
    result = parse_resume_bytes(_make_docx_bytes(), "resume.docx")
    assert result.file_type == "DOCX"
    assert "Python" in result.text
    assert "Flask 系统" in result.text


def test_parse_pdf_text_and_page_count() -> None:
    result = parse_resume_bytes(_make_pdf_bytes(), "resume.pdf")
    assert result.file_type == "PDF"
    assert "Python Flask SQL resume" in result.text
    assert result.page_count == 1


@pytest.mark.parametrize("content", [b"", b"  \r\n\t"])
def test_reject_empty_text(content: bytes) -> None:
    with pytest.raises(ResumeParseError):
        parse_txt_bytes(content)


def test_reject_too_long_text() -> None:
    with pytest.raises(ResumeParseError, match="超过"):
        parse_txt_bytes("测试内容".encode("utf-8"), max_characters=3)


def test_reject_oversized_txt() -> None:
    with pytest.raises(ResumeParseError, match="2 MB"):
        parse_txt_bytes(b"a" * (MAX_TXT_BYTES + 1))


def test_reject_oversized_document() -> None:
    with pytest.raises(ResumeParseError, match="10 MB"):
        parse_resume_bytes(b"a" * (MAX_DOCUMENT_BYTES + 1), "resume.pdf")


def test_reject_unknown_text_encoding() -> None:
    with pytest.raises(ResumeParseError, match="编码无法识别"):
        parse_txt_bytes(b"\x81")


def test_reject_unsupported_extension() -> None:
    with pytest.raises(ResumeParseError, match="仅支持"):
        parse_resume_bytes(b"legacy document", "resume.doc")


def test_reject_fake_pdf() -> None:
    with pytest.raises(ResumeParseError, match="不是有效的 PDF"):
        parse_resume_bytes(b"not a pdf", "resume.pdf")


def test_reject_fake_docx() -> None:
    with pytest.raises(ResumeParseError, match="损坏|格式异常"):
        parse_resume_bytes(b"PK broken docx", "resume.docx")


def test_empty_pdf_mentions_ocr() -> None:
    with pytest.raises(ResumeParseError, match="OCR"):
        parse_resume_bytes(_make_pdf_bytes(text=None), "scanned.pdf")


def test_reject_password_protected_pdf() -> None:
    document = fitz.open()
    document.new_page().insert_text((72, 72), "Protected resume")
    protected = document.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner-password",
        user_pw="user-password",
    )
    document.close()

    with pytest.raises(ResumeParseError, match="已加密"):
        parse_resume_bytes(protected, "protected.pdf")
